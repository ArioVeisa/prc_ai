import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

doc = docx.Document("/home/ray/Downloads/Template_Laporan_Praktikum_Individu_Kecerdasan Artifisial dan Penerapannya.docx")

# 1. Update Cover Title & Info
doc.paragraphs[2].text = "Studi Kasus Mandiri: Klasifikasi Risiko Penyakit Jantung Menggunakan Machine Learning Library (Random Forest) dan AI Model Service (LLM Few-Shot) Berbasis Streamlit Deployment"
doc.paragraphs[3].text = "Laporan Praktikum – Studi Kasus Mandiri (Heart Failure Prediction)"
doc.paragraphs[5].text = "Laporan Praktikum — Studi Kasus Mandiri (Heart Failure Prediction)"

# Clear instruction box paragraph 6
doc.paragraphs[6].text = ""

# Table 0: Metadata
meta_table = doc.tables[0]
meta_table.rows[0].cells[1].text = "Ario veisa Rayanda Utomo"
meta_table.rows[1].cells[1].text = "1204230086"
meta_table.rows[2].cells[1].text = "SI-06-REG"
meta_table.rows[3].cells[1].text = "Kecerdasan Artifisial dan Penerapannya"
meta_table.rows[4].cells[1].text = "Ganjil 2026–2027"
meta_table.rows[5].cells[1].text = "28/10/2026"
meta_table.rows[6].cells[1].text = "M. Nizar P. Ma’ady, S.Kom., M.IM., M.Kom."

# 2. Section 1: Tujuan Praktikum
doc.paragraphs[8].text = "Tujuan pelaksanaan praktikum mandiri ini adalah:"
doc.paragraphs[9].text = "1. Membangun dan mengevaluasi model klasifikasi data tabular menggunakan Machine Learning Library (Random Forest Pipeline) untuk memprediksi risiko penyakit jantung secara akurat dan efisien."
doc.paragraphs[10].text = "2. Merancang strategi rekayasa prompt dan in-context few-shot learning pada AI Model Service (Large Language Model) untuk klasifikasi risiko klinis berbasis penalaran semantik fitur medis."
doc.paragraphs[11].text = "3. Mengintegrasikan model library dan AI model service ke dalam antarmuka web Streamlit interaktif (input formulir pasien, prediksi batch CSV, dan fitur unduh contoh dataset) serta men-deploy aplikasi ke Streamlit Community Cloud."

# 3. Section 2: Dasar Teori Singkat
doc.paragraphs[13].text = (
    "Machine Learning Library (khususnya ensemble tree seperti Random Forest) merupakan standar industri dalam pemodelan data tabular terstruktur. "
    "Random Forest bekerja dengan membangun sekumpulan decision tree independen dan menggabungkan hasil prediksinya melalui mekanisme majority voting. "
    "Dengan bantuan modul ColumnTransformer, fitur numerik diskalakan menggunakan StandardScaler sedangkan fitur kategori dikonversi melalui OneHotEncoder. "
    "Karakteristik utama pendekatan library lokal ini adalah kecepatan inferensi sangat tinggi (< 5 ms), deterministik, hemat komputasi, dan tidak memerlukan koneksi internet saat runtime.\n\n"
    "Sebaliknya, AI Model Service berbasis Large Language Model (LLM) memanfaatkan kapabilitas in-context learning. "
    "Tanpa memperbarui bobot parameter internal (no retraining), model diberikan instruksi sistemik (system prompt) dan beberapa contoh representatif (few-shot demonstrations) dalam konteks input. "
    "LLM mampu memahami relasi semantik klinis antar-gejala (misalnya hubungan antara nyeri dada asimptomatik/ASY dengan depresi segmen ST latihan/Oldpeak). "
    "Namun, pendekatan service memiliki konsekuensi latensi tinggi akibat network round-trip time serta biaya token API per panggilan.\n\n"
    "Streamlit berperan sebagai layer deployment reaktif berbasis Python yang menyajikan model machine learning lokal (.pkl) dan remote API service ke dalam antarmuka web yang mudah digunakan. "
    "Dengan fitur manajemen rahasia (st.secrets) dan caching data (@st.cache_data), aplikasi web dapat berjalan aman dan hemat kuota di Streamlit Community Cloud."
)

# 4. Section 3: Alat dan Bahan
doc.paragraphs[15].text = "Perangkat keras: PC / Laptop x86_64, CPU 8-Core, RAM 16 GB, SSD 512 GB."
doc.paragraphs[16].text = "Sistem operasi: Linux Ubuntu 24.04 LTS (x86_64)."
doc.paragraphs[17].text = "Perangkat lunak: Python 3.12, scikit-learn 1.9.1, pandas 3.0.6, joblib 1.6.0, openai 1.66.1, streamlit 1.66.1, git 2.43.0."
doc.paragraphs[18].text = (
    "Akun/layanan & data: Repositori GitHub (https://github.com/ArioVeisa/prc_ai), Streamlit Community Cloud, "
    "AI Model Service API (OpenAI Compatible endpoint), dan Heart Failure Prediction Dataset (Kaggle/UCI, 918 baris, 11 fitur klinis)."
)

# 5. Section 4: Hasil & Pembahasan

# --- Langkah 1 ---
doc.paragraphs[20].text = "Langkah 1 — Persiapan Data & Pemodelan Machine Learning Library (Random Forest)"
doc.paragraphs[21].text = (
    "Dataset Heart Failure Prediction (918 sampel) dibagi menjadi data latih 80% (734 sampel) dan data uji 20% (184 sampel) "
    "secara terstratifikasi berdasarkan label target (Normal vs Penyakit Jantung). "
    "Sebuah pipeline scikit-learn dibangun menggunakan ColumnTransformer untuk menstandarisasi 6 fitur numerik dan melakukan one-hot encoding pada 5 fitur kategori. "
    "Model dilatih dengan RandomForestClassifier(n_estimators=100, max_depth=6, random_state=42), dievaluasi pada data uji, "
    "lalu disimpan sebagai artefak model_heart.pkl beserta model_info.json."
)

# Table 1: Code Langkah 1
t1_cell = doc.tables[1].cell(0, 0)
t1_cell.text = (
    "# Potongan kode training Random Forest Pipeline (train_model.py)\n"
    "preprocessor = ColumnTransformer(transformers=[\n"
    "    ('num', StandardScaler(), ['Age', 'RestingBP', 'Cholesterol', 'FastingBS', 'MaxHR', 'Oldpeak']),\n"
    "    ('cat', OneHotEncoder(handle_unknown='ignore'), ['Sex', 'ChestPainType', 'RestingECG', 'ExerciseAngina', 'ST_Slope'])\n"
    "])\n"
    "model = Pipeline([\n"
    "    ('preprocessor', preprocessor),\n"
    "    ('classifier', RandomForestClassifier(n_estimators=100, max_depth=6, random_state=42))\n"
    "])\n"
    "model.fit(train_df[FEATURES], train_df['Target'])\n"
    "y_pred = model.predict(test_df[FEATURES])\n"
    "joblib.dump(model, 'model_heart.pkl')"
)
for p in t1_cell.paragraphs:
    p.style = doc.styles["Normal"]
    for r in p.runs:
        r.font.name = "Courier New"
        r.font.size = Pt(8.5)

# Table 2: Image Langkah 1
t2_cell = doc.tables[2].cell(0, 0)
t2_cell.text = ""
p2 = t2_cell.paragraphs[0]
p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
p2.add_run().add_picture("studi_kasus_mandiri/figures/confusion_matrices_heart.png", width=Inches(5.6))

doc.paragraphs[25].text = "Gambar 1. Matriks Konfusi dan Evaluasi Model Random Forest vs LLM pada Data Uji."
doc.paragraphs[26].text = (
    "Analisis: Pada pengujian data uji (N=184), model Random Forest berhasil meraih akurasi sebesar 88.59% dan Macro F1-Score sebesar 0.8837. "
    "Pada evaluasi sampel uji acak 50 pasien (Gambar 1 kiri), model mencatatkan 20 True Normal, 25 True Penyakit Jantung, 3 False Positive, dan 2 False Negative "
    "(Akurasi 90.0%, F1 0.8990). Kecepatan inferensi model library sangat unggul, yakni hanya 3.08 milidetik per sampel, "
    "menjadikannya opsi ideal untuk pemrosesan data bervolume besar secara instan."
)

# --- Langkah 2 ---
doc.paragraphs[27].text = "Langkah 2 — Rekayasa Prompt & Evaluasi AI Model Service (LLM Few-Shot)"
doc.paragraphs[28].text = (
    "Pada tahap ini, klasifikasi dilakukan menggunakan Large Language Model Service (ai-practicum) via API. "
    "Konfigurasi disimpan dalam config_heart.json dengan menyusun system prompt yang menjelaskan konteks klinis kardiologi, "
    "indikator risiko mayor (seperti ChestPainType ASY, ST_Slope Flat, dan Oldpeak tinggi), serta 6 contoh data few-shot representatif (3 Normal, 3 Penyakit Jantung). "
    "Fungsi classify() diimplementasikan dalam classifier_heart.py dengan mekanisme pemanggilan OpenAI-compatible API dan parsing label deterministik."
)

# Table 3: Code Langkah 2
t3_cell = doc.tables[3].cell(0, 0)
t3_cell.text = (
    "# Potongan logika prompt & classifier LLM Service (classifier_heart.py)\n"
    "def build_messages(cfg, data_input):\n"
    "    messages = [{'role': 'system', 'content': cfg['system_prompt']}]\n"
    "    for fs in cfg.get('few_shot', []):\n"
    "        messages.append({'role': 'user', 'content': format_data(fs['data'])})\n"
    "        messages.append({'role': 'assistant', 'content': fs['label']})\n"
    "    messages.append({'role': 'user', 'content': format_data(data_input)})\n"
    "    return messages\n\n"
    "def classify(client, cfg, data_input):\n"
    "    resp = client.chat.completions.create(\n"
    "        model=cfg['model'], temperature=0.0, max_tokens=10,\n"
    "        messages=build_messages(cfg, data_input)\n"
    "    )\n"
    "    return parse_label(resp.choices[0].message.content, cfg['labels'])"
)
for p in t3_cell.paragraphs:
    p.style = doc.styles["Normal"]
    for r in p.runs:
        r.font.name = "Courier New"
        r.font.size = Pt(8.5)

# Table 4: Image Langkah 2
t4_cell = doc.tables[4].cell(0, 0)
t4_cell.text = ""
p4 = t4_cell.paragraphs[0]
p4.alignment = WD_ALIGN_PARAGRAPH.CENTER
p4.add_run().add_picture("studi_kasus_mandiri/figures/metrics_comparison_heart.png", width=Inches(5.4))

doc.paragraphs[32].text = "Gambar 2. Grafik Perbandingan Metrik Evaluasi: Random Forest vs LLM Service."
doc.paragraphs[33].text = (
    "Analisis: Evaluasi 50 data uji menunjukkan bahwa LLM Service mencapai akurasi sebesar 94.0% dan Macro F1-Score sebesar 0.9398 (Gambar 2). "
    "LLM berhasil menekan False Positive menjadi hanya 1 kasus (Gambar 1 kanan) berkat kemampuannya membedakan profil klinis asimptomatik dengan depresi ST. "
    "Meskipun akurasi LLM lebih tinggi 4% dibanding Random Forest pada sampel uji, latensinya mencapai 9.132,80 ms (sekitar 9,1 detik per sampel) "
    "atau 2.965 kali lebih lambat dibanding Random Forest. Hal ini menegaskan adanya trade-off mendasar antara fleksibilitas penalaran LLM dan efisiensi waktu komputasi library."
)

# --- Langkah 3 ---
doc.paragraphs[34].text = "Langkah 3 — Pembuatan Web Streamlit, Fitur Unduh Sampel, dan Deployment Cloud"
doc.paragraphs[35].text = (
    "Aplikasi web CardioGuard AI dibangun menggunakan Streamlit (app.py) dengan antarmuka terpadu yang memadukan kedua model. "
    "Aplikasi memiliki formulir input manual 11 fitur klinis (Tab 1) dan fasilitas batch prediction file CSV (Tab 2). "
    "Sesuai instruksi penugasan, tombol unduh sampel data uji (sample_test_heart.csv) disediakan langsung di sidebar aplikasi. "
    "Proyek kemudian di-push ke GitHub dan di-deploy ke Streamlit Community Cloud dengan menyematkan API key via fitur Secrets."
)

# Table 5: Code Langkah 3
t5_cell = doc.tables[5].cell(0, 0)
t5_cell.text = (
    "# Potongan antarmuka Streamlit (studi_kasus_mandiri/app.py)\n"
    "with st.sidebar:\n"
    "    st.header('❤️ CardioGuard AI')\n"
    "    engine = st.radio('Pilih Model:', ['Komparasi Keduanya', 'Random Forest', 'LLM Service'])\n"
    "    st.download_button('Unduh sample_test_heart.csv', data=sample_csv_bytes, file_name='sample_test_heart.csv')\n\n"
    "tab1, tab2 = st.tabs(['🩺 Prediksi Satu Pasien', '📁 Prediksi Batch dari File CSV'])\n"
    "with tab1:\n"
    "    # Form 11 input klinis pasien...\n"
    "    if st.form_submit_button('Analisis Risiko Pasien'):\n"
    "        # Inferensi Random Forest (proba) & LLM Service\n"
    "        st.success(f'Hasil: {hasil_prediksi}')"
)
for p in t5_cell.paragraphs:
    p.style = doc.styles["Normal"]
    for r in p.runs:
        r.font.name = "Courier New"
        r.font.size = Pt(8.5)

# Table 6: Image Langkah 3
t6_cell = doc.tables[6].cell(0, 0)
t6_cell.text = ""
p6 = t6_cell.paragraphs[0]
p6.alignment = WD_ALIGN_PARAGRAPH.CENTER
p6.add_run().add_picture("studi_kasus_mandiri/figures/latency_comparison_heart.png", width=Inches(4.6))

doc.paragraphs[39].text = "Gambar 3. Grafik Perbandingan Latensi Eksekusi Inferensi (Skala Logaritmik)."
doc.paragraphs[40].text = (
    "Analisis: Antarmuka web berhasil memfasilitasi interaksi pengguna secara komprehensif. Pengguna dapat memilih model yang diinginkan atau menjalankan komparasi konsensus AI. "
    "Fitur unduh contoh data memungkinkan penguji langsung mengunduh 15 sampel pasien terverifikasi lalu mengunggahnya ke tab CSV tanpa perlu menyiapkan format manual. "
    "Aplikasi telah aktif dan ter-deploy secara publik pada URL Streamlit Community Cloud:\n"
    "URL Deployment: https://arioveisa-prc-ai-studi-kasus-mandiriapp-0nybuc.streamlit.app\n"
    "Repositori GitHub: https://github.com/ArioVeisa/prc_ai/tree/main/studi_kasus_mandiri"
)

# 6. Section 6: Kendala dan Solusi
k_table = doc.tables[7]
# Row 1
k_table.rows[1].cells[1].text = "FileNotFoundError saat memuat model (.pkl), metadata, dan konfigurasi (.json) saat aplikasi pertama kali dijalankan di Streamlit Community Cloud."
k_table.rows[1].cells[2].text = "Menggunakan modul Path(__file__).resolve().parent untuk menentukan direktori kerja dinamis file aplikasi, sehingga pencarian aset model selalu valid terlepas dari root direktori kerja server cloud."

# Row 2
k_table.rows[2].cells[1].text = "Konflik pembacaan environment secrets di Streamlit Cloud ketika kunci GROQ_API_KEY terdefinisi sebagai string kosong, menyebabkan fallback ke endpoint praktikum terblokir."
k_table.rows[2].cells[2].text = "Memperbaiki alur percabangan pembacaan secret dengan menerapkan fungsi .strip() dan validasi non-empty sebelum instansiasi client API, serta menyelaraskan isian menu Secrets pada dashboard Streamlit Cloud."

# Row 3
k_table.rows[3].cells[1].text = "Latensi pemanggilan AI Model Service yang mencapai ~9 detik per data menyebabkan proses prediksi batch CSV terasa membeku (unresponsive)."
k_table.rows[3].cells[2].text = "Menerapkan dekorator @st.cache_data pada fungsi prediksi LLM, menambahkan progress bar dinamis saat iterasi baris, dan membatasi inferensi interaktif LLM pada 15 baris sampel teratas."

# 7. Section 7: Kesimpulan
doc.paragraphs[45].text = (
    "Berdasarkan pelaksanaan studi kasus mandiri klasifikasi risiko penyakit jantung (Heart Failure Prediction) "
    "menggunakan Machine Learning Library (Random Forest) dan AI Model Service (LLM Few-Shot) berbasis Streamlit Deployment, "
    "dapat disimpulkan bahwa:\n"
    "1. Pemodelan Machine Learning Library (Random Forest Pipeline) terbukti sangat efektif untuk data tabular terstruktur, "
    "menghasilkan akurasi 88.59% pada keseluruhan data uji (dan 90.0% pada evaluasi sampel) dengan keunggulan kecepatan inferensi superior sebesar 3.08 ms per sampel tanpa membutuhkan koneksi internet.\n"
    "2. Pemodelan AI Model Service (LLM Few-Shot) menunjukkan kapabilitas penalaran klinis semantik yang tinggi dengan akurasi 94.0% dan Macro F1-Score 0.9398. "
    "LLM mampu mengidentifikasi korelasi kompleks antar-fitur gejala tanpa retraining, namun memiliki kendala latensi tinggi (9.132,80 ms/sampel) dan biaya kuota token.\n"
    "3. Seluruh instruksi penugasan praktikum berhasil dipenuhi secara tuntas: integrasi form input manual otomatis, tab prediksi batch CSV, "
    "penyediaan file data sampel siap unduh (sample_test_heart.csv), serta keberhasilan deployment publik pada URL: "
    "https://arioveisa-prc-ai-studi-kasus-mandiriapp-0nybuc.streamlit.app."
)

output_docx = "studi_kasus_mandiri/Laporan_Praktikum_Individu_Heart_Disease.docx"
doc.save(output_docx)
print(f"Laporan docx berhasil dibuat: {output_docx}")
